from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def extract_text(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = file_path.suffix.lower()

    if ext == ".pdf":
        reader = PdfReader(str(file_path))
        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)

        return "\n\n".join(pages).strip()

    elif ext == ".docx":
        doc = Document(str(file_path))
        return "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ).strip()

    elif ext == ".pptx":
        presentation = Presentation(str(file_path))
        slides = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slides.append(shape.text)

        return "\n\n".join(slides).strip()

    elif ext == ".txt":
        return file_path.read_text(encoding="utf-8").strip()

    else:
        raise ValueError(f"Unsupported file type: {ext}")


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python document_parser.py <file>")
        sys.exit(1)

    text = extract_text(sys.argv[1])

    print("\n========== EXTRACTED TEXT ==========\n")
    print(text)
